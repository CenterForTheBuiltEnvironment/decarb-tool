"""One-time extraction of tests/testing_cases.xlsx into a frozen JSON fixture.

The spreadsheet itself is git-ignored and kept locally; the committed JSON is
what the tests read.

    python -m tests.verification.extract_spreadsheet

The spreadsheet holds hand calculations made independently of the code for
7 scenarios x 3 hours on 3 buildings. This script freezes BOTH its expected
values AND the equipment data it assumed into
tests/fixtures/spreadsheet_cases.json, so later edits to the production
equipment library cannot break the comparison. Rerun only if the spreadsheet
itself is corrected; review the JSON diff before committing.

Where the spreadsheet's inputs differ from today's library, the spreadsheet
wins (it is the frozen reference). Known differences, applied as overrides:
  * hr01 COP at 32.2 °C, first point: 6.7 (library now 5.7)
  * hr01 refrigerant 200 kg, GWP 466 (library now 400 kg, GWP 292)
  * hr03 refrigerant 29 kg (library now 45 kg)
  * chillers: COP 2.84, GWP 2088, 703.2 kW / 90.7 kg (B1S3: 351.6 kW / 45.4 kg)

Values the spreadsheet computed with a rule the tool has since replaced are
kept in the fixture but tagged "superseded" with the reason, and are not
asserted:
  * AWHP cooling in hours where the AWHPs are also heating. The spreadsheet
    gives such hours no AWHP cooling at all; the tool now lets the compressors
    not needed for heating serve cooling (verified separately in
    test_verification_handcalc.py::test_awhp_cooling_compressor_split).
"""

import copy
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "tests" / "testing_cases.xlsx"
LIBRARY = ROOT / "data" / "input" / "equipment_data.JSON"
OUT = ROOT / "tests" / "fixtures" / "spreadsheet_cases.json"

SCENARIO_DEFAULTS = {
    "hr_wwhp": None,
    "hr_wwhp_performance_model": "interpolate_HHWST",
    "hr_wwhp_h_supply_t": None,
    "awhp": None,
    "awhp_performance_model": "interpolate_HHWST_fixed",
    "awhp_h_supply_t": None,
    "awhp_sizing_mode": "integer_sizing_peak_load",
    "awhp_sizing_value": 1,
    "awhp_redundancy": 1,
    "awhp_use_cooling": False,
    "awhp_sizing_priority": "heating",
    "fuel_switching": False,
}

# Transcribed from the "Scenarios" sheet (equipment IDs and settings per case).
CASES = {
    "B1S1": dict(hr_wwhp="hr03", hr_wwhp_h_supply_t=48.9, awhp="hp01", awhp_h_supply_t=38,
                 awhp_use_cooling=True, backup_heating="res02", chiller="xs_chiller_703kW"),
    "B1S2": dict(awhp="hp01", awhp_performance_model="interpolate_HHWST_reset", awhp_h_supply_t=38,
                 awhp_redundancy=2, awhp_use_cooling=True, backup_heating="res01",
                 chiller="xs_chiller_703kW"),
    "B1S3": dict(hr_wwhp="hr03", hr_wwhp_h_supply_t=52, awhp="hp01", awhp_h_supply_t=42,
                 awhp_sizing_mode="fractional_sizing_peak_load", awhp_sizing_value=0.7,
                 awhp_use_cooling=True, backup_heating="bo03", chiller="xs_chiller_352kW"),
    "B2S1": dict(backup_heating="res02", chiller="xs_chiller_703kW"),
    "B2S2": dict(hr_wwhp="hr03", hr_wwhp_h_supply_t=60, awhp="hp06", awhp_h_supply_t=60,
                 awhp_sizing_mode="fixed_num_units", awhp_sizing_value=3, backup_heating="res02",
                 chiller="xs_chiller_703kW"),
    "B3S1": dict(hr_wwhp="hr01", hr_wwhp_h_supply_t=40, awhp="hp01", awhp_h_supply_t=38,
                 awhp_sizing_mode="fractional_sizing_peak_load", awhp_sizing_value=0.3,
                 awhp_use_cooling=True, backup_heating="bo02", chiller="xs_chiller_703kW"),
    "B3S2": dict(awhp="hp04", awhp_h_supply_t=45, awhp_sizing_value=0.5, awhp_use_cooling=True,
                 backup_heating="res02", chiller="xs_chiller_703kW"),
}  # fmt: skip

# (sheet section, sheet row label) -> (tool column, divisor from tool units to sheet units)
ENERGY_CHECKS = {
    ("HR WWHP", "HR HHW Output [kW]"): ("hr_hhw_W", 1000),
    ("HR WWHP", "HR CHW Output [kW]"): ("hr_chw_W", 1000),
    ("HR WWHP", "HR-WWHP Electricity [kWh]"): ("elec_hr_Wh", 1000),
    ("AWHP Heating", "AWHP Count (Heating)"): ("awhp_num", 1),
    ("AWHP Heating", "AWHP Heating COP"): ("awhp_cop_h", 1),
    ("AWHP Heating", "AWHP Heating Cap [kW]"): ("awhp_cap_h_W", 1000),
    ("AWHP Heating", "AWHP HHW Output [kW]"): ("awhp_hhw_W", 1000),
    ("AWHP Heating", "AWHP Heating Electricity [kWh]"): ("elec_awhp_h_Wh", 1000),
    ("Boiler", "Total Gas [kWh]"): ("gas_Wh", 1000),
    ("Res. Heat", "Resistance Heater Electricity [kWh]"): ("elec_res_Wh", 1000),
    ("AWHP Cooling", "AWHP Cooling COP"): ("awhp_cop_c", 1),
    ("AWHP Cooling", "AWHP Cooling Cap [kW]"): ("awhp_cap_c_W", 1000),
    ("AWHP Cooling", "AWHP CHW Output [kW]"): ("awhp_chw_W", 1000),
    ("AWHP Cooling", "AWHP Cooling Electricity [kWh]"): ("elec_awhp_c_Wh", 1000),
    ("Chiller", "Chiller CHW Output [kW]"): ("chiller_chw_W", 1000),
    ("Chiller", "Chiller Electricity [kWh]"): ("elec_chiller_Wh", 1000),
    ("Chiller", "Total Electricity [kWh]"): ("elec_Wh", 1000),
}
AWHP_COOLING_DEPENDENT = {
    "awhp_cap_c_W", "awhp_chw_W", "elec_awhp_c_Wh", "chiller_chw_W", "elec_chiller_Wh", "elec_Wh",
}  # fmt: skip
# Genuine, still-open method differences. The test runs these as strict xfails,
# so a code change that resolves one turns the suite red until this list is updated.
KNOWN_DIFFERENCES = {
    ("B1S3", 0, "awhp_num"): (
        "AWHP sizing reference capacity at HHWST 42 °C: spreadsheet interpolates the "
        "38 and 52 °C tables (176.35 kW at 0 °C); the tool uses the nearest table, 38 °C "
        "(176.92 kW), giving 11.098 instead of 11.134 units (-0.3 %). The spreadsheet's "
        "other two hours already use the tool's 11.1."
    ),
}
KNOWN_DIFFERENCES[("B1S3", 0, "awhp_cap_h_W")] = KNOWN_DIFFERENCES[("B1S3", 0, "awhp_num")]
SUPERSEDED_REASON = (
    "spreadsheet predates the AWHP compressor split: it allows no AWHP cooling in hours "
    "where the AWHPs also heat"
)


def _read_sheet(wb, name):
    """{(section, label): [21 values]} — 7 cases x 3 hours in columns D..X."""
    rows = list(wb[name].iter_rows(values_only=True))
    out, section = {}, None
    for r in rows[4:]:
        section = r[1] or section
        if r[2]:
            out[(section, r[2])] = list(r[3:24])
    return out, list(rows[3][3:24])


def _chiller(eq_id, capacity_W, refrigerant_kg):
    return {
        "eq_id": eq_id,
        "eq_type": "chiller",
        "eq_calc_type": "specific",
        "model": f"Spreadsheet chiller {capacity_W / 1000:.0f} kW",
        "fuel": "electricity",
        "capacity_W": capacity_W,
        "refrigerant": "R-spreadsheet",
        "refrigerant_weight_g": refrigerant_kg * 1000,
        "refrigerant_gwp": 2088,
        "performance": {"cooling": {"efficiency": 2.84}},
    }


def _equipment():
    library = {e["eq_id"]: e for e in json.loads(LIBRARY.read_text())["equipment"]}
    used = {v for case in CASES.values() for k, v in case.items() if k in
            ("hr_wwhp", "awhp", "backup_heating") and v}  # fmt: skip
    eq = {k: copy.deepcopy(library[k]) for k in sorted(used)}
    eq["hr01"]["performance"]["heating"]["leaving_supply_t"]["32.2"]["cop"][0] = 6.7
    eq["hr01"]["refrigerant_weight_g"], eq["hr01"]["refrigerant_gwp"] = 200_000, 466
    eq["hr03"]["refrigerant_weight_g"] = 29_000
    eq["xs_chiller_703kW"] = _chiller("xs_chiller_703kW", 703_200, 90.7)
    eq["xs_chiller_352kW"] = _chiller("xs_chiller_352kW", 351_600, 45.4)
    return list(eq.values())


def extract() -> dict:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    energy, timestamps = _read_sheet(wb, "Energy")
    emis, _ = _read_sheet(wb, "Emissions")
    scen_rows = {r[1]: r[2:9] for r in wb["Scenarios"].iter_rows(values_only=True) if r[1]}

    cases = []
    for i, (case_id, overrides) in enumerate(CASES.items()):
        scenario = {
            "eq_scen_id": case_id,
            "eq_scen_name": case_id,
            **SCENARIO_DEFAULTS,
            **overrides,
        }
        cols = range(3 * i, 3 * i + 3)
        awhp_cooling = bool(scenario["awhp"]) and scenario["awhp_use_cooling"]
        peak_heat = energy[("AWHP Heating", "Peak HHW load (kW)")][cols[0]]

        hours = []
        for hour_idx, c in enumerate(cols):
            awhp_heating_now = (energy[("AWHP Heating", "AWHP HHW Output [kW]")][c] or 0) > 0
            cooling_now = (energy[("Load", "Cooling Load [kW]")][c] or 0) > 0
            # With no cooling load only the reported cooling capacity differs; with a
            # cooling load the whole AWHP-cooling / chiller split differs.
            superseded_cols = (
                (AWHP_COOLING_DEPENDENT if cooling_now else {"awhp_cap_c_W"})
                if awhp_cooling and awhp_heating_now
                else set()
            )
            checks, superseded, known = [], [], []
            for (section, label), (column, divisor) in ENERGY_CHECKS.items():
                value = energy.get((section, label), [None] * 21)[c]
                if value is None or isinstance(value, str):
                    continue
                if section == "AWHP Cooling" and not awhp_cooling:
                    continue  # not applicable: spreadsheet lists values for absent equipment
                entry = {"label": label, "column": column, "divisor": divisor, "value": value}
                if (case_id, hour_idx, column) in KNOWN_DIFFERENCES:
                    known.append(
                        {**entry, "reason": KNOWN_DIFFERENCES[(case_id, hour_idx, column)]}
                    )
                elif column in superseded_cols:
                    superseded.append({**entry, "reason": SUPERSEDED_REASON})
                else:
                    checks.append(entry)
            hours.append(
                {
                    "timestamp": timestamps[c].isoformat(),
                    "t_out_C": energy[("Load", "Outdoor Temp [°C]")][c],
                    "heating_kW": energy[("Load", "Heating Load [kW]")][c],
                    "cooling_kW": energy[("Load", "Cooling Load [kW]")][c],
                    "checks": checks,
                    "superseded": superseded,
                    "known_differences": known,
                    "elec_rate_g_per_kWh": emis[
                        ("Electricity", "Elec Emissions Rate [g CO2e/kWh]")
                    ][c],
                    "elec_emissions_kg": emis[("Electricity", "Electricity Emissions [kg CO2e]")][
                        c
                    ],
                    "gas_emissions_kg": emis[("Gas", "Gas Emissions [kg CO2e]")][c],
                }
            )

        j = list(CASES).index(case_id)
        cases.append(
            {
                "id": case_id,
                "scenario": scenario,
                "peak_heating_kW": peak_heat,
                "peak_chiller_kW": emis[("Refrigerant", "Peak chiller output (kW)")][cols[0]],
                "emissions": {
                    "grid_scenario": scen_rows["Grid Scenario"][j],
                    "gea_grid_region": scen_rows["GEA Grid Region"][j],
                    "emission_type": scen_rows["Emission Type"][j],
                    "shortrun_weighting": scen_rows["Short-run weighting"][j],
                    "year": scen_rows["Year"][j],
                    "annual_refrig_leakage": emis[("Refrigerant", "Annual leakage")][cols[0]],
                    "ng_rate_g_per_kWh": emis[("Gas", "Gas emissions rate (g/kWh)")][cols[0]],
                    "refrig_emissions_kg_per_hour": emis[
                        ("Refrigerant", "Refrigerant Emissions [kg CO2e]")
                    ][cols[0]],
                },
                "hours": hours,
            }
        )

    return {
        "source": "tests/testing_cases.xlsx",
        "generated_by": "python -m tests.verification.extract_spreadsheet",
        "equipment": _equipment(),
        "cases": cases,
    }


if __name__ == "__main__":
    OUT.write_text(json.dumps(extract(), indent=1, default=str) + "\n")
    print(f"Wrote {OUT.relative_to(ROOT)}")
